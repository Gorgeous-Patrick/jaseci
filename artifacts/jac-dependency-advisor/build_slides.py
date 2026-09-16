"""Build an editable PowerPoint, a vector PDF, speaker notes, and previews.

Dependencies: python-pptx, reportlab, pymupdf, pillow.
Run with standard CPython: python build_slides.py
"""
from __future__ import annotations

import json
import math
from pathlib import Path
import shutil
import subprocess

from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
import pymupdf as fitz
from PIL import Image, ImageDraw, ImageFont

OUT = Path(__file__).resolve().parent
STEM = 'jac_dependency_idea'
W, H = 960, 540
C = {
    'paper': '#F7F8F5', 'white': '#FFFFFF', 'ink': '#142C3B',
    'muted': '#576C78', 'line': '#D7E0E1', 'navy': '#122B3A',
    'teal': '#008577', 'teal_light': '#E0F2EC', 'blue': '#3563B2',
    'blue_light': '#E8EEFA', 'amber': '#9B620F', 'amber_light': '#FBEDD3',
    'red': '#AF4E43', 'red_light': '#F8E6E1', 'soft': '#EBEFEC',
    'night_text': '#D4E2E4', 'mint': '#7CDAC6',
}
SOURCES = {
    'inspection': 'https://onlinelibrary.wiley.com/doi/abs/10.1002/cpe.4330030607',
    'galois': 'https://www.cs.cornell.edu/~kb/publications/Galois_PLDI07.pdf',
    'boosting': 'https://erickoskinen.com/papers/boosting-ppopp08.pdf',
    'legion': 'https://theory.stanford.edu/~aiken/publications/papers/sc12.pdf',
    'pr': 'https://github.com/jaseci-labs/jac/pull/9215',
    'commit': '3c3a7ffb2eaafb1083219dde9a9fb7583c301346',
}


def font_path(query):
    return subprocess.check_output(['fc-match', '-f', '%{file}', query], text=True)


FONT_REG = font_path('Liberation Sans')
FONT_BOLD = font_path('Liberation Sans:style=Bold')
FONT_MONO = font_path('DejaVu Sans Mono')
for name, path in [('Deck', FONT_REG), ('DeckBold', FONT_BOLD), ('DeckMono', FONT_MONO)]:
    pdfmetrics.registerFont(TTFont(name, path))


def wrap(text, width, size, font):
    result = []
    for paragraph in text.split('\n'):
        if not paragraph:
            result.append('')
            continue
        line = ''
        for word in paragraph.split(' '):
            candidate = (line + ' ' + word) if line else word
            if line and pdfmetrics.stringWidth(candidate, font, size) > width:
                result.append(line)
                line = word
            else:
                line = candidate
        result.append(line)
    return result


class Slide:
    def __init__(self, title, section, dark=False, appendix=False):
        self.title, self.section = title, section
        self.dark, self.appendix = dark, appendix
        self.elements = []
        self.notes = ''

    def box(self, x, y, w, h, fill='white', stroke=None, radius=0):
        self.elements.append(dict(kind='box', x=x, y=y, w=w, h=h,
                                  fill=C.get(fill, fill), stroke=C.get(stroke, stroke), radius=radius))

    def line(self, x1, y1, x2, y2, color='line', width=1.5, arrow=False, dashed=False):
        self.elements.append(dict(kind='line', x1=x1, y1=y1, x2=x2, y2=y2,
                                  color=C.get(color, color), width=width, arrow=arrow, dashed=dashed))

    def text(self, text, x, y, w, size=19, color='ink', bold=False,
             mono=False, align='left', max_h=None, leading=1.19, link=None):
        font = 'DeckMono' if mono else ('DeckBold' if bold else 'Deck')
        lines = wrap(text, w, size, font)
        height = len(lines) * size * leading + 4
        if max_h is not None and height > max_h:
            raise ValueError(f'Slide {self.title}: text too tall: {text!r} ({height} > {max_h})')
        if x < 0 or y < 0 or x+w > W+0.01 or y+height > H-5:
            raise ValueError(f'Slide {self.title}: text outside bounds: {text!r}')
        for line in lines:
            if pdfmetrics.stringWidth(line, font, size) > w+0.5:
                raise ValueError(f'Text line too wide: {line!r}')
        self.elements.append(dict(kind='text', lines=lines, x=x, y=y, w=w, h=height,
                                  size=size, color=C.get(color, color), bold=bold, mono=mono,
                                  font=font, align=align, leading=leading, link=link))
        return height

    def node(self, label, x, y, r=26, fill='teal_light', stroke='teal', text_color='ink'):
        self.elements.append(dict(kind='circle', x=x-r, y=y-r, w=r*2, h=r*2,
                                  fill=C.get(fill, fill), stroke=C.get(stroke, stroke)))
        self.text(label, x-r, y-11, 2*r, size=17, bold=True, color=text_color, align='center')

    def edge(self, x1, y1, x2, y2, r1=26, r2=26, color='teal', dashed=False):
        a = math.atan2(y2-y1, x2-x1)
        self.line(x1+math.cos(a)*(r1+2), y1+math.sin(a)*(r1+2),
                  x2-math.cos(a)*(r2+6), y2-math.sin(a)*(r2+6),
                  color, 2.1, arrow=True, dashed=dashed)

    def header(self, subtitle=None):
        self.text(self.section.upper(), 52, 28, 850, 10, 'mint' if self.dark else 'teal', True)
        self.text(self.title, 52, 59, 856, 32, 'white' if self.dark else 'ink', True,
                  max_h=88, leading=1.1)
        if subtitle:
            self.text(subtitle, 52, 116, 856, 17, 'night_text' if self.dark else 'muted', max_h=47)

    def footer(self, number):
        self.line(52, 503, 908, 503, '#34505D' if self.dark else 'line', 0.7)
        label = 'JAC  /  QUERY-DERIVED DEPENDENCIES'
        if self.appendix:
            label += '  /  APPENDIX'
        self.text(label, 52, 513, 770, 8.5, 'night_text' if self.dark else 'muted')
        self.text(f'{number:02d}', 873, 510, 35, 11, 'mint' if self.dark else 'teal', True, align='right')


slides = []

# 1. Thesis. The three-stage picture previews the logic of the talk.
s = Slide('Resolve Dependencies\nBefore Abilities Run', 'Research proposal', dark=True)
s.text('RESEARCH PROPOSAL  /  SEPTEMBER 2026', 52, 31, 840, 11, 'mint', True)
s.text(s.title, 52, 90, 850, 47, 'white', True, max_h=126, leading=1.07)
s.text('Use Jac graph queries to decide which abilities can start.', 54, 217, 850, 23, 'night_text')
for x, no, heading, body in [
    (52, '01', 'Extract accesses', 'The compiler identifies queries and possible writes.'),
    (349, '02', 'Find the nodes', 'Before the body runs, inspect the current graph.'),
    (646, '03', 'Schedule abilities', 'Wait for earlier conflicting work. Start independent work.'),
]:
    s.box(x, 295, 262, 116, '#1C3A49')
    s.text(no, x+18, 309, 36, 12, 'mint', True)
    s.text(heading, x+18, 336, 228, 20, 'white', True)
    s.text(body, x+18, 368, 228, 15, 'night_text')
s.line(317, 354, 339, 354, 'mint', 2, True)
s.line(614, 354, 636, 354, 'mint', 2, True)
s.text('Can a small amount of inspection prevent a large amount of wasted computation?',
       52, 442, 850, 21, 'white', max_h=53)
s.notes = '''The proposal is to inspect dependencies before starting an ability body. The compiler recognizes graph queries and summarizes possible reads and writes. For an eligible queued invocation, the scheduler already knows its current node, walker, and arguments. It can run a separate, side-effect-free inspection function to resolve those accesses to actual node identities in the current graph. It then compares pending invocations and starts independent work while holding back calls that depend on earlier work. The ability itself has not started during inspection. This differs from waiting for a running body to reach its query. The graph does not have to be known at compile time: the compiler provides the inspection code, and the current runtime graph supplies the identities. Eligibility requires known query inputs and safe inspection. Summaries may be conservative, and dependencies must be refreshed when their inputs change. Existing speculative validation and ordered effects remain the correctness baseline. This is an unimplemented research proposal, not a measured speedup. Eight slides develop the argument; two appendices cover prior work and the existing prototype.'''
slides.append(s)

# Shared scaffold: two abilities, one node, explicit values, left-to-right time.
# P02 shows the wasted attempt; P06 changes only when B does its calculation.
def value_example(s):
    s.box(52, 151, 856, 74, 'white', 'line')
    s.text('STARTING VALUE', 70, 164, 173, 10, 'muted', True)
    s.text('n.score = 10', 70, 187, 178, 20, 'ink', mono=True)
    s.line(265, 164, 265, 212, 'line', 1)
    s.text('WHAT THE SERIAL PROGRAM REQUIRES', 285, 164, 598, 10, 'muted', True)
    s.text('A sets score to 20. B must calculate using 20.', 285, 187, 598, 20, 'ink', True)
    s.text('TIME', 170, 244, 56, 10, 'muted', True)
    s.line(219, 250, 905, 250, 'muted', 1.2, True)
    s.text('Writer A', 52, 293, 107, 17, 'blue', True)
    s.text('Reader B', 52, 380, 107, 17, 'ink', True)
    s.box(170, 279, 185, 55, 'blue_light')
    s.text('Prepare 20', 180, 297, 165, 18, 'blue', True, align='center')
    s.line(363, 308, 559, 308, 'blue', 1.3, True, dashed=True)
    s.text('Not visible yet', 372, 282, 179, 14, 'muted', align='center')
    s.box(573, 279, 134, 55, 'blue_light')
    s.text('Publish 20', 580, 297, 120, 18, 'blue', True, align='center')
    s.line(640, 336, 640, 356, 'blue', 1.7, True)


s = Slide('B can finish a calculation using the wrong input.', '01 / Problem')
s.header('Two abilities share a node n. The original program runs A before B.')
value_example(s)
s.box(170, 362, 121, 60, 'amber_light')
s.text('Read 10', 178, 381, 105, 18, 'amber', True, align='center')
s.line(296, 392, 307, 392, 'muted', 1.3, True)
s.box(313, 362, 244, 60, 'red_light')
s.text('Calculate using 10', 325, 381, 220, 20, 'red', True, align='center')
s.line(561, 392, 570, 392, 'muted', 1.3, True)
s.box(573, 362, 134, 60, 'red_light')
s.text('Discard\nresult', 581, 373, 118, 17, 'red', True, align='center')
s.line(712, 392, 741, 392, 'muted', 1.3, True)
s.box(747, 362, 161, 60, 'teal_light')
s.text('Recalculate\nusing 20', 755, 373, 145, 17, 'teal', True, align='center')
s.line(170, 431, 170, 437, 'red', 1.4)
s.line(170, 437, 557, 437, 'red', 1.4)
s.line(557, 431, 557, 437, 'red', 1.4)
s.text('First attempt: wasted computation', 170, 446, 388, 15, 'red', True, align='center')
s.text('The input was 10. The required input was 20. B has to calculate again.',
       52, 476, 856, 19, 'ink', True)
s.notes = '''Use the explicit numbers before discussing speculation. There is one shared node n, initially with score 10. The original serial program runs ability A first: A changes the score to 20. Ability B then reads the score and performs an expensive calculation. So the correct input to B is 20. Follow the two rows from left to right. A prepares its update, but that speculative update is still private. B reads the currently visible value, 10, and calculates using 10. Only after the speculative work completes does A publish 20 in serial commit order. When the runtime checks B's attempt, it finds that B used the old value. It discards that result and runs B again with 20. The red bracket identifies precisely the wasted attempt. Publishing means making the buffered update visible. The diagram illustrates event order and values; horizontal lengths are not measured durations. It describes the current prototype's ordered commit-time validation, not a claim that all transactional memory systems detect conflicts late. Slide 6 uses this exact example to show how inspecting dependencies before launching B can avoid the first calculation.'''
slides.append(s)

# 3. Clearly separate compiler extraction, runtime inspection, and the ability body.
s = Slide('Run a small inspection before the ability body.', '02 / Idea')
s.header('For queries whose starting node and parameters are already known.')
s.box(52, 169, 403, 283, 'navy')
s.text('IN THE COMPILER', 73, 189, 359, 11, 'mint', True)
s.text('Recognize B’s graph query:', 73, 224, 359, 21, 'white', True)
s.text('[->:Edge:->]', 73, 266, 359, 31, 'white', mono=True)
s.text('Generate a separate function that\nfinds B’s possible read/write targets.',
       73, 326, 359, 20, 'white', max_h=78)
s.text('The ability body stays intact.', 73, 419, 359, 16, 'mint', True)
s.line(465, 303, 485, 303, 'teal', 2.4, True)
s.box(495, 169, 413, 283, 'white', 'line')
s.text('IN THE SCHEDULER — BEFORE B STARTS', 515, 189, 370, 11, 'teal', True)
s.text('B’s current node is known.', 515, 225, 370, 21, 'ink', True)
s.text('Inspect its outgoing Edge links:', 515, 267, 370, 19)
s.box(515, 309, 372, 77, 'teal_light')
s.text('The query finds n1 and n2.', 533, 324, 336, 21, 'teal', True)
s.text('B may read these nodes.', 533, 358, 336, 16)
s.text('B has not started its computation.', 515, 419, 371, 16, 'teal', True)
s.text('Jac’s built-in query tells the compiler how to find the relevant nodes.',
       52, 474, 856, 19, 'ink', True)
s.notes = '''This slide explains why Jac helps. Its built-in query explicitly provides a graph operation, an origin, a direction, and an edge type. The compiler can use this known meaning to generate a separate inspection function. The scheduler runs that function on the actual current graph before launching the ability body. For this example, B's here and query parameters are already known, so inspection finds n1 and n2. If B reads their fields, these become conservative whole-node read targets. Simply obtaining a reference does not inherently read every field: whole-node access is an initial approximation, and the compiler must also analyze how the result is used. Assignments and calls need corresponding conservative write summaries. The inspection itself has no user-visible effects; the body stays intact and is not split into a computation phase and an effect phase. Only safe, available inputs may be inspected early. Queries depending on arbitrary earlier body computation, own writes, or unknown effectful predicates need conservative treatment or the existing speculative/serial fallback. The compiler knows how to look up the nodes; only runtime inspection knows their current identities. An OOP library with equivalent semantic summaries could support the same strategy, but Jac provides a standard construct from which to derive it automatically.'''
slides.append(s)

# 4. Match concrete pending calls, including potential writes, before launching them.
s = Slide('Match reads with writes before choosing what starts.', '03 / Scheduling')
s.header('Three queued ability calls. Their original serial order is A, then B, then C.')
s.box(52, 170, 504, 281, 'white', 'line')
s.text('INSPECTED NODE ACCESSES', 74, 190, 458, 11, 'teal', True)
s.text('Call', 74, 233, 110, 16, 'muted', True)
s.text('May read', 225, 233, 165, 16, 'muted', True)
s.text('May write', 412, 233, 121, 16, 'muted', True)
for yy, label, reads, writes, fill in [
    (267, 'A', '—', 'n2', 'blue_light'),
    (311, 'B', 'n1, n2', '—', 'teal_light'),
    (355, 'C', '—', 'n9', 'soft'),
]:
    s.box(67, yy, 474, 39, fill)
    s.text(label, 81, yy+8, 100, 20, 'ink', True)
    s.text(reads, 225, yy+9, 164, 19)
    s.text(writes, 412, yy+9, 119, 19)
s.text('B needs n2. A will change n2 first.', 74, 416, 458, 18, 'ink', True)
s.box(581, 170, 327, 281, 'white', 'line')
s.text('START A AND C TOGETHER', 603, 190, 283, 11, 'teal', True)
s.box(603, 237, 125, 51, 'blue_light')
s.text('A', 613, 249, 105, 23, 'blue', True, align='center')
s.box(759, 237, 125, 51, 'teal_light')
s.text('C', 769, 249, 105, 23, 'teal', True, align='center')
s.line(665, 294, 665, 349, 'blue', 2, True)
s.text('A publishes', 687, 310, 191, 17, 'blue')
s.box(603, 355, 125, 51, 'blue_light')
s.text('B', 613, 367, 105, 23, 'blue', True, align='center')
s.text('Then start B', 743, 371, 143, 17, 'ink', True)
s.text('C’s effects still publish after B.', 603, 425, 283, 14, 'muted')
s.text('Keep the original order for conflicting calls. Independent calls can compute together.',
       52, 474, 856, 18, bold=True)
s.notes = '''These are three concrete queued invocations, not every ability definition in the program. An invocation includes its current node, walker, and arguments. Inspection finds that A may write n2, B may read n1 and n2, and C may write an unrelated n9. Assume no other conflicting shared state or topology changes in this example. The original order is A, B, C. Therefore B must wait for A to publish its update. A and C can compute concurrently, but C's buffered effects are still published after B. A query result alone is insufficient: the compiler must supply conservative write summaries as well, including aliases and called functions. For any pair i before j, a conservative scheduler orders them when a possible write in either call overlaps a read or write in the other. This preserves read-after-write, write-after-read, and write-after-write relationships; reads alone do not conflict. Never move a later writer before an earlier reader merely because the reader accesses its node. Use an index from stable object identities to pending readers and writers if needed; the implementation data structure is not the research claim. Query inputs, adjacency, predicate reads, and non-node state must also be covered. New calls created by visit can be inspected when the runtime has identified them; no full knowledge of future traversal is assumed.'''
slides.append(s)

# 5. The phantom dependency case makes the correctness condition tangible.
s = Slide('If earlier work changes the query, inspect again.', '04 / Correctness')
s.header('The node set found before execution is valid only while the query’s inputs stay unchanged.')
for x, label in [(52, 'INSPECTION FINDS NO MATCHES'), (498, 'AN EARLIER ABILITY ADDS AN EDGE')]:
    s.box(x, 173, 410, 193, 'white', 'line')
    s.text(label, x+20, 191, 370, 11, 'muted', True)
s.node('here', 155, 273, 29, 'blue_light', 'blue')
s.text('Result: []', 222, 260, 214, 20, 'muted', mono=True)
s.edge(588, 273, 805, 273)
s.node('here', 588, 273, 29, 'blue_light', 'blue')
s.node('n4', 805, 273, 29)
s.text('Edge', 674, 243, 72, 14, 'teal')
s.text('Result: [n4]', 626, 323, 244, 19, 'teal', mono=True)
s.box(52, 388, 856, 56, 'amber_light')
s.text('Before B starts: query again and update its dependencies to include n4.',
       70, 406, 820, 21, 'amber', True)
s.text('Track changes to the searched edges and filter inputs, even when the result is empty.',
       52, 465, 856, 17, 'muted')
s.notes = '''Pre-execution inspection does not make a query result permanently valid. Suppose B's inspection finds no outgoing Edge links, but an earlier ability A adds a link to n4. B must use the updated graph in the serial program. Tracking only the originally returned nodes would miss this change because the old result was empty. Record the graph data examined at every hop, including intermediate nodes' searched adjacency and empty lookups. Conservative adjacency versions can detect insertions, deletions, and ordering changes. Also cover node and edge liveness, type matching, and all predicate inputs, including fields of rejected candidates. Intermediate query results can feed later hops inside the inspection function; the ability body need not start. When earlier work invalidates an inspection, re-run it and update scheduling dependencies before launching the body. The same rule applies if a possible writer's target set changes. Re-inspection must happen against a coherent view, and its dependencies must be registered atomically with validation or protected by an equivalent protocol; a bare version check followed by an unprotected launch is insufficient. Earlier calls whose accesses cannot be bounded retain conservative barriers or speculative validation. Keep the existing instrumentation initially, so inspection improves scheduling without being the sole correctness mechanism. The current prototype serializes graph mutation through fallback. A first performance experiment can keep topology fixed and exercise mutation/invalidation correctness separately.'''
slides.append(s)

# 6. Same values as slide 2; B remains queued until the known dependency clears.
s = Slide('B waits before starting, then calculates once.', '05 / Expected benefit')
s.header('Before either body starts, inspection finds: A writes n; B reads n.')
value_example(s)
s.box(170, 362, 387, 60, 'soft')
s.text('B stays queued — its body has not started',
       183, 383, 361, 17, 'muted', True, align='center')
s.line(561, 392, 570, 392, 'muted', 1.3, True)
s.box(573, 362, 134, 60, 'blue_light')
s.text('Start B\nRead 20', 581, 373, 118, 17, 'blue', True, align='center')
s.line(712, 392, 741, 392, 'muted', 1.3, True)
s.box(747, 362, 161, 60, 'teal_light')
s.text('Calculate once\nusing 20', 755, 373, 145, 17, 'teal', True, align='center')
s.text('No calculation with 10. No discarded first attempt.',
       170, 441, 738, 18, 'teal', True)
s.text('Refresh inspection if its inputs changed. Keep final validation for unproven accesses.',
       52, 478, 856, 15, 'muted')
s.notes = '''Return to the example from slide 2: n.score starts at 10, A precedes B, and B must calculate using 20. This time the compiler-generated inspection runs before either ability body starts. It finds a possible write by A to n and a possible read by B from n. The scheduler leaves B in the queue while A computes and publishes 20. B has not executed a prefix of its body and has not occupied a worker waiting at the query. After the dependency is satisfied and the inspection is validated or refreshed as necessary, B starts, reads 20, and calculates once. The red wasted attempt from slide 2 never occurs in this example. A node-field update does not necessarily change which nodes a topology-only query returns; only inspection inputs that changed require re-inspection. The body still reads the current field value after A publishes. This schedule is qualitative, not a timing measurement. Retain final validation for unknown or incompletely summarized accesses. Implementing this schedule requires dependency-aware task admission and a way to commit ready earlier work and then release dependents; the current batch-join-before-commit loop cannot simply launch a waiting dependent in that same batch. The research claim is about reduced total work after paying for inspection, not a guaranteed speedup on every workload.'''
slides.append(s)

# 7. The trade-off is the research question; avoid overclaiming speedup.
s = Slide('The hypothesis is a measurable trade-off.', '06 / Evaluation')
s.header('Inspection is useful when avoiding wasted computation saves more time than inspection costs.')
s.box(52, 172, 412, 160, 'teal_light')
s.text('Potential benefit', 74, 192, 367, 23, 'teal', True)
s.text('Fewer failed speculative attempts\nWorkers start independent calls first', 74, 235, 366, 19, max_h=83)
s.box(496, 172, 412, 160, 'amber_light')
s.text('Potential cost', 518, 192, 367, 23, 'amber', True)
s.text('Extra lookups and re-inspection\nConservative sets may delay safe work', 518, 235, 366, 19, max_h=83)
s.text('CONTROLLED COMPARISON', 52, 353, 270, 10, 'muted', True)
cols = [(52, 206, 'Serial'), (270, 206, 'Current speculation'),
        (488, 206, 'Eager memory tracking'), (706, 202, 'Inspect then schedule')]
for x, ww, label in cols:
    s.box(x, 374, ww, 51, 'white', 'line')
    s.text(label, x+10, 391, ww-20, 15, 'ink', True, align='center')
s.text('Vary: shared nodes, query size, computation cost, and graph changes.', 52, 445, 856, 17, bold=True)
s.text('Measure: total time, discarded computation, inspection cost, and unnecessary waits.',
       52, 473, 856, 16, 'muted')
s.notes = '''The hypothesis is that inexpensive pre-execution inspection can save more work than it adds. Measure both sides: inspection performs graph lookups and dependency comparisons, may duplicate queries later executed by the body, and may need to run again after relevant changes. Conservative whole-node or may-write sets can also delay calls that would not actually conflict. Compare serial execution, current speculation, eager memory conflict tracking, and inspect-then-schedule under the same worker count, task boundaries, memory management, and original-order effect publication. An eager memory baseline helps separate the effect of earlier conflict detection from the specific benefit of compiler-generated graph inspection. Initially keep ordinary tracking and validation in the proposed system. Vary shared-neighbor overlap, query fan-out, computation cost, and graph mutation rate; start performance experiments with fixed topology to isolate the mechanism. Measure end-to-end time, discarded CPU work, inspector and scheduler overhead, and delays caused by conservative summaries. Fewer retries alone do not establish success if inspection or lost parallelism costs more than it saves. Reusing a query result is an optional later optimization requiring a proof that its inputs and observable order remain valid.'''
slides.append(s)

# 8. Research decision; stays centered on the advisor conversation.
s = Slide('Start with one query and one falsifiable claim.', '07 / Proposed next step', dark=True)
s.header('Can finding dependencies before launch reduce total execution time?')
items = [
    (52, '01', 'Inspect complete queries', 'Available, safe inputs.\nInclude multi-hop queries.\nSummarize reads and writes.'),
    (347, '02', 'Schedule before launch', 'Hold back conflicting calls.\nKeep ordered effects.\nRetain speculative validation.'),
    (642, '03', 'Measure the full cost', 'Check identical behavior.\nInclude inspection overhead.\nFind when it pays off.'),
]
for x, no, heading, body in items:
    s.box(x, 177, 266, 215, '#1C3A49')
    s.text(no, x+18, 196, 225, 15, 'mint', True)
    s.text(heading, x+18, 232, 229, 23, 'white', True, max_h=65)
    s.text(body, x+18, 307, 228, 16, 'night_text', max_h=76)
s.text('The contribution to establish', 52, 427, 856, 13, 'mint', True)
s.text('Use graph syntax to generate dependency inspection,\nthen schedule ability calls before their bodies run.',
       52, 447, 856, 23, 'white', True, leading=1.08, max_h=57)
s.notes = '''Support complete eligible graph queries, including multi-hop chains. Eligibility depends on whether query inputs can be obtained safely before the ability body starts, not on hop count. A later hop can consume an earlier hop's result entirely within the inspector. The native compiler already lowers query chains through refs and refs_many; inspector generation can reuse this query structure, while adding conservative read/write summaries and tracking graph data examined at every hop. Generate inspectors for queued native entry-ability calls and schedule before their bodies start. Preserve original-order effects and existing speculative validation. Unknown calls, accesses whose targets depend on unavailable body results, and unsupported effects keep conservative or existing fallback handling; the experiment does not require splitting general ability bodies into computation and effect phases. Check shared neighbors, aliases, read-before-write order, multiple writers, empty intermediate results, intermediate-edge changes, and predicate invalidation. Future calls produced by visit are inspected once they are known. For performance, vary hop count, fan-out, overlap, and computation cost, accounting for every inspection. The candidate research contribution is automatically generating useful dependency inspection from Jac graph syntax and integrating it with ordered walker semantics. Inspector-executor scheduling and semantic concurrency control already exist. The contribution must be established by extraction coverage, soundness, and a measured advantage over appropriate baselines, not by the scheduling pattern alone.'''
slides.append(s)

# 9. Verified primary references; explicit epistemic status of novelty.
s = Slide('Inspection before execution has clear precedents.', 'Appendix A / Positioning', appendix=True)
s.header('The candidate contribution is automatic extraction from Jac, with its ordered walker behavior.')
rows = [
    ('Runtime preprocessing', 'Saltz, Berryman & Wu · 1991',
     'Compiler-generated runtime preprocessing can help expose parallel work.', SOURCES['inspection']),
    ('Galois', 'Kulkarni et al. · PLDI 2007',
     'Abstractions and runtime conflict detection expose parallelism in irregular programs.', SOURCES['galois']),
    ('Legion', 'Bauer et al. · SC 2012',
     'Logical regions and access privileges provide dependence information for task scheduling.', SOURCES['legion']),
]
for y, (name, cite, desc, url) in zip([175, 266, 357], rows):
    s.line(52, y-7, 908, y-7, 'line', 1)
    s.text(name, 52, y+4, 328, 21, 'ink', True, link=url)
    s.text(cite, 52, y+39, 342, 12, 'muted', link=url)
    s.text(desc, 417, y+7, 486, 18, max_h=70)
s.box(52, 455, 856, 36, 'teal_light')
s.text('Inspect first, execute second. Study what Jac can extract automatically, and at what cost.',
       63, 466, 836, 15, 'teal', True)
s.notes = '''The proposal belongs near inspector-executor approaches: a preliminary runtime step examines accesses or structure, and a later step performs the actual work according to that information. Saltz, Berryman, and Wu describe the value of runtime preprocessing and its integration into compilers; this is a precedent for the overall direction, not evidence for the proposed Jac implementation. Galois establishes the role of abstractions in optimistic parallelism for irregular object and graph programs. Legion uses regions and declared access privileges to expose task dependencies. Transactional boosting, listed in sources.md as additional context, uses abstract operations and commutativity for transactional objects. An OOP implementation can support similar ideas. The candidate distinction here is automatically extracting inspectors from Jac's existing query syntax, resolving them for concrete queued walker calls before body execution, and preserving the original order of effects. This is a candidate contribution rather than an established novelty claim. A deeper related-work review and a direct evaluation are still necessary. Each displayed reference links to a primary source.'''
slides.append(s)

# 10. Existing evidence is labeled separately from the proposal.
bench_src = Path('/tmp/jac-walker-speedup.json')
bench_dst = OUT / 'prototype_benchmark.json'
if bench_src.exists():
    shutil.copyfile(bench_src, bench_dst)
if not bench_dst.exists():
    raise FileNotFoundError('The original prototype benchmark JSON is required for appendix B.')
bench = json.loads(bench_dst.read_text())
s = Slide('The current prototype is the experimental baseline.', 'Appendix B / Evidence and status', appendix=True)
s.header('Existing native walker speculation provides an execution substrate and a serial reference.')
s.text('Implemented', 52, 174, 373, 24, 'teal', True)
s.text('Opt-in speculative entry abilities\nBuffered effects and ordered commit\nSerial fallback for unsupported effects',
       52, 217, 380, 20, max_h=112)
s.text('Proposed', 52, 359, 373, 24, 'blue', True)
s.text('Pre-execution query inspection\nand dependency-aware scheduling',
       52, 402, 380, 20, max_h=72)
s.box(475, 172, 433, 263, 'white', 'line')
s.text('EXISTING CPU MICROBENCHMARK', 495, 191, 390, 10, 'muted', True)
s.text('Mode', 496, 226, 155, 14, 'muted', True)
s.text('Median', 661, 226, 89, 14, 'muted', True, align='right')
s.text('Speedup', 773, 226, 112, 14, 'muted', True, align='right')
for yy, key, label in [(260, 'off', 'Serial / off'), (309, 'auto2', '2 threads'), (358, 'auto4', '4 threads')]:
    s.line(495, yy-10, 887, yy-10, 'line', 0.8)
    accent = 'teal' if key == 'auto4' else 'ink'
    s.text(label, 496, yy, 153, 20, accent, key == 'auto4')
    s.text(f"{bench['median_seconds'][key]*1000:.0f} ms", 652, yy, 103, 20, accent, key == 'auto4', align='right')
    s.text(f"{bench['speedup'][key]:.2f}×", 775, yy, 110, 20, accent, True, align='right')
s.text('16 independent tasks; identical output; 0 fallbacks.', 495, 408, 390, 12, 'muted')
s.text('Baseline measurement: 2 warm-ups, 7 timed runs; native AOT + RC; 16 logical CPUs.',
       52, 456, 856, 14, 'muted')
s.text('Pre-execution inspection remains unevaluated. Source: PR #9215 and bundled benchmark JSON.',
       52, 480, 856, 13, 'muted', link=SOURCES['pr'])
s.notes = '''The existing prototype is opt-in and defaults to off. It supports native AOT programs with reference counting on supported 64-bit Linux targets. It speculates across entry abilities, buffers supported effects, validates reads at ordered commit, and falls back to serial execution for unsupported operations such as graph mutation. Exit execution remains serial. The reported benchmark is the previously collected synthetic conflict-free CPU workload, not a graph-dependency experiment: 16 jobs with 6.4 million floating-point steps per job, two warm-ups and seven timed samples. Median wall times were approximately 424 ms with speculation off, 341 ms with two worker threads, and 223 ms with four, giving 1.90x at four threads. Output matched and there were 16 commits and zero fallbacks. It used native AOT, RC, optimization level 2, host libc, and a host with 16 logical CPUs. The measurements were made with the local toolchain compatibility setup described in the PR. They support feasibility of parallel execution for this synthetic kernel, not a general speedup or any claim about query-derived dependency tracking. Full samples are bundled in prototype_benchmark.json. Source checkout for this deck: commit 3c3a7ffb2.'''
slides.append(s)


def rgb(s):
    return RGBColor.from_string(s.lstrip('#'))


def render():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(W), Pt(H)
    prs.core_properties.title = 'Resolve Dependencies Before Abilities Run'
    prs.core_properties.subject = 'Research proposal: pre-execution graph inspection and scheduling in Jac'
    prs.core_properties.author = ''
    prs.core_properties.keywords = 'Jac, graph queries, speculative execution, runtime dependencies'
    prs.core_properties.comments = '8 main slides plus 2 appendix slides. Proposed mechanisms are explicitly identified.'
    pdf_path = OUT / f'{STEM}.pdf'
    cv = canvas.Canvas(str(pdf_path), pagesize=(W, H), pageCompression=1)
    cv.setTitle(prs.core_properties.title)
    cv.setAuthor('')
    for num, s in enumerate(slides, 1):
        s.footer(num)
        ps = prs.slides.add_slide(prs.slide_layouts[6])
        ps.background.fill.solid()
        bg = C['navy'] if s.dark else C['paper']
        ps.background.fill.fore_color.rgb = rgb(bg)
        cv.setFillColor(HexColor(bg))
        cv.rect(0, 0, W, H, fill=1, stroke=0)
        cv.bookmarkPage(f'slide-{num}')
        cv.addOutlineEntry(s.title.replace('\n', ' '), f'slide-{num}', level=0)
        for e in s.elements:
            kind = e['kind']
            if kind in ('box', 'circle'):
                shape_type = MSO_AUTO_SHAPE_TYPE.OVAL if kind == 'circle' else MSO_AUTO_SHAPE_TYPE.RECTANGLE
                shape = ps.shapes.add_shape(shape_type, Pt(e['x']), Pt(e['y']), Pt(e['w']), Pt(e['h']))
                shape.fill.solid()
                shape.fill.fore_color.rgb = rgb(e['fill'])
                if e.get('stroke'):
                    shape.line.color.rgb = rgb(e['stroke'])
                    shape.line.width = Pt(1)
                else:
                    shape.line.fill.background()
                cv.setFillColor(HexColor(e['fill']))
                cv.setStrokeColor(HexColor(e.get('stroke') or e['fill']))
                cv.setLineWidth(1)
                if kind == 'circle':
                    cv.ellipse(e['x'], H-e['y']-e['h'], e['x']+e['w'], H-e['y'], fill=1, stroke=int(bool(e.get('stroke'))))
                else:
                    cv.rect(e['x'], H-e['y']-e['h'], e['w'], e['h'], fill=1, stroke=int(bool(e.get('stroke'))))
            elif kind == 'line':
                shape = ps.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(e['x1']), Pt(e['y1']), Pt(e['x2']), Pt(e['y2']))
                shape.line.color.rgb = rgb(e['color'])
                shape.line.width = Pt(e['width'])
                if e['dashed']:
                    from pptx.enum.dml import MSO_LINE_DASH_STYLE
                    shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
                if e['arrow']:
                    arrow = OxmlElement('a:tailEnd')
                    arrow.set('type', 'triangle')
                    arrow.set('w', 'med')
                    arrow.set('len', 'med')
                    shape.line._get_or_add_ln().append(arrow)
                cv.setStrokeColor(HexColor(e['color']))
                cv.setLineWidth(e['width'])
                cv.setDash([4, 4] if e['dashed'] else [])
                cv.line(e['x1'], H-e['y1'], e['x2'], H-e['y2'])
                cv.setDash([])
                if e['arrow']:
                    angle = math.atan2(e['y2']-e['y1'], e['x2']-e['x1'])
                    pts = [(e['x2'], e['y2'])]
                    for sign in [-1, 1]:
                        pts.append((e['x2'] - 8*math.cos(angle) + sign*3.5*math.sin(angle),
                                    e['y2'] - 8*math.sin(angle) - sign*3.5*math.cos(angle)))
                    path = cv.beginPath()
                    path.moveTo(pts[0][0], H-pts[0][1])
                    for px, py in pts[1:]:
                        path.lineTo(px, H-py)
                    path.close()
                    cv.setFillColor(HexColor(e['color']))
                    cv.drawPath(path, stroke=0, fill=1)
            elif kind == 'text':
                shape = ps.shapes.add_textbox(Pt(e['x']), Pt(e['y']), Pt(e['w']), Pt(e['h']))
                tf = shape.text_frame
                tf.word_wrap = False
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Pt(0)
                tf.vertical_anchor = MSO_ANCHOR.TOP
                for i, line in enumerate(e['lines']):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[e['align']]
                    p.space_before = p.space_after = Pt(0)
                    p.line_spacing = Pt(e['size']*e['leading'])
                    r = p.add_run()
                    r.text = line
                    r.font.name = 'Consolas' if e['mono'] else 'Arial'
                    r.font.size = Pt(e['size'])
                    r.font.bold = e['bold']
                    r.font.color.rgb = rgb(e['color'])
                    if e.get('link'):
                        r.hyperlink.address = e['link']
                cv.setFont(e['font'], e['size'])
                cv.setFillColor(HexColor(e['color']))
                ascent = pdfmetrics.getAscent(e['font']) / 1000 * e['size']
                for i, line in enumerate(e['lines']):
                    yy = H-e['y']-ascent-i*e['size']*e['leading']
                    xx = e['x']
                    if e['align'] == 'center':
                        xx += e['w']/2
                        cv.drawCentredString(xx, yy, line)
                    elif e['align'] == 'right':
                        xx += e['w']
                        cv.drawRightString(xx, yy, line)
                    else:
                        cv.drawString(xx, yy, line)
                if e.get('link'):
                    cv.linkURL(e['link'], (e['x'], H-e['y']-e['h'], e['x']+e['w'], H-e['y']), relative=0, thickness=0)
        ps.notes_slide.notes_text_frame.text = s.notes
        cv.showPage()
    cv.save()
    prs.save(OUT / f'{STEM}.pptx')

    notes = ['# Resolve Dependencies Before Abilities Run', '',
             'English advisor deck. Slides 1–8 form the main argument; slides 9–10 are optional appendices.', '',
             '## One-sentence thesis', '',
             'Generate inspectors from graph queries and read/write summaries, resolve dependencies before launching eligible ability bodies, and test whether scheduling from those dependencies reduces total work while preserving serial behavior.', '']
    for i, s in enumerate(slides, 1):
        notes += [f'## {i:02d}. {s.title.replace(chr(10), " ")}', '', s.notes, '']
    (OUT / 'speaker_notes.md').write_text('\n'.join(notes), encoding='utf-8')
    (OUT / 'slide_text.json').write_text(json.dumps([
        dict(title=s.title, section=s.section,
             text=['\n'.join(e['lines']) for e in s.elements if e['kind']=='text'], notes=s.notes)
        for s in slides], indent=2, ensure_ascii=False), encoding='utf-8')

    # Render the actual PDF for visual verification and a convenient overview.
    doc = fitz.open(pdf_path)
    preview_dir = OUT / 'previews'
    preview_dir.mkdir(exist_ok=True)
    tiles = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        page_path = preview_dir / f'slide-{i+1:02d}.png'
        pix.save(page_path)
        image = Image.open(page_path).convert('RGB')
        image.thumbnail((576, 324), Image.Resampling.LANCZOS)
        tiles.append(image.copy())
    overview = Image.new('RGB', (1188, 5*354+24), '#E1E6E4')
    draw = ImageDraw.Draw(overview)
    font = ImageFont.truetype(FONT_REG, 13)
    for i, tile in enumerate(tiles):
        x, y = 12+(i%2)*588, 12+(i//2)*354
        overview.paste(tile, (x, y))
        draw.text((x+4, y+329), f'{i+1:02d}  {slides[i].title.replace(chr(10), " ")}', font=font, fill=C['ink'])
    overview.save(OUT / 'slide_overview.png')
    assert len(doc) == 10
    for i, page in enumerate(doc):
        assert page.get_text().strip(), f'Empty PDF page {i+1}'
        for block in page.get_text('dict')['blocks']:
            if 'lines' not in block:
                continue
            for line in block['lines']:
                for span in line['spans']:
                    x0, y0, x1, y1 = span['bbox']
                    assert x0 >= 0 and y0 >= 0 and x1 <= W+1 and y1 <= H+1, (i, span)
    loaded = Presentation(OUT / f'{STEM}.pptx')
    assert len(loaded.slides) == 10
    assert all(len(x.notes_slide.notes_text_frame.text) > 100 for x in loaded.slides)
    print(f'Created and validated {len(slides)} slides in {OUT}')
    print('PDF and PowerPoint: 16:9; diagrams and text are vectors/editable shapes.')


if __name__ == '__main__':
    render()
